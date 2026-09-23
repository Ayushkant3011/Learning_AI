import os
from pathlib import Path
from dotenv import load_dotenv
from groq import Groq
import time

load_dotenv()
my_api_key = os.getenv("GROQ_API_KEY")

if not my_api_key:
    raise ValueError("API KEY NOT FOUND !!!!")

client = Groq(api_key = my_api_key)
model="openai/gpt-oss-20b"

job_description="""
    Description
At Audible, we believe stories have the power to transform lives. It's why we work with some of the world's leading creators to produce and share audio storytelling with our millions of global listeners. We are dreamers and inventors who come from a wide range of backgrounds and experiences to empower and inspire each other. Imagine your future with us.

ABOUT YOU
If you are a dedicated problem solver, this is the role for you. Identifying and overcoming challenges is at the heart of everything you do; whether its looking ahead to our listeners problems, or analyzing every inefficiency, you thrive when you strive for perfection. We look for people who take responsibility, who feel the pride in a job well done and see every failure as an opportunity to learn. With a knowledgeable team around you, and opportunities to challenge yourself, its a perfect environment to develop and apply your new skills.

As a Software Development Engineer, you will...
- Develop, test, and maintain software components and applications for the world's largest audiobook creator and provider
- Identify and resolve software defects through code review, log analysis, metrics examination, debugging, and other troubleshooting methods
- Write maintainable and scalable code for software components as part of an agile engineering team
- Collaborate with stakeholders to envision, design, develop, test, and launch customer-centric software that inspires and impacts users
- Apply computer science, engineering, and analysis principles to solve moderate-scale problems
- Continuously enhance testability, operational excellence (OE), and documentation for owned software components
- Handle data classification, storage, and management in accordance with company policies
- Participate in the interview process and provide mentorship to fellow engineers

ABOUT AUDIBLE
Audible is the leading producer and provider of audio storytelling. We spark listeners' imaginations, offering immersive, cinematic experiences full of inspiration and insight to enrich our customers' daily lives. We are a global company with an entrepreneurial spirit. We are dreamers and inventors who are passionate about the positive impact Audible can make for our customers and our neighbors. This spirit courses throughout Audible, supporting a culture of creativity and inclusion built on our People Principles and our mission to build more equitable communities in the cities we call home.

Basic Qualifications
- Experience programming with at least one software programming language
- Bachelor's degree or equivalent in Computer Science

Preferred Qualifications
- Experience from previous technical internship(s) or demonstrated project experience
- Experience with Cloud platforms (preferably AWS), database systems (SQL and NoSQL), AI tools for development productivity, contributing to open-source projects, and/or version control systems
- Experience delivering large, cross-functional, customer facing products
"""
from pydantic import BaseModel
class JobDesc(BaseModel):
    role:str
    required_skills:list[str]
    preferred_skills:list[str]
    minimum_experience: float | None
    education_requirements:list[str]
    responsibilities:list[str]


jobDesc_Schema = JobDesc.model_json_schema()

system_prompt=f"""
    You are an expert HR Assistant.
    
    Your job is to analyze the job description and extract structured information from them.

    Return ONLY valid JSON matching this schema:{jobDesc_Schema}

    IMPORTANT:
    Do NOT return the schema itself
    Do NOT return fields like "properties", "title", or "type".
    Fill the schema with the actual information extracted from the job description.

    If minimum experience is not mentioned, return null.
    If Information for a list is missing, return an empty list.
    Do not invent information.
"""


user_prompt=f"""
    Analyze the following job description: 
    
    {job_description}
"""


message_system={
    "role":"system",
    "content": system_prompt
}

message_user={
    "role":"user",
    "content": user_prompt
}

response_format={
    "type":"json_object"
}


messages=[message_system, message_user]

response = client.chat.completions.create(model=model, messages = messages, response_format = response_format)

answer = response.choices[0].message.content

raw_json = answer

import json
job_data= json.loads(raw_json)

job = JobDesc(**job_data)

print(job.minimum_experience)
print(job.education_requirements)



# Parse real
class MatchResult(BaseModel):
    score: float
    details: dict

class Experience(BaseModel):
    company: str | None = None
    role:str | None = None
    duration:str | None = None
    description:str | None = None
    skills_used:str | None = None

class Resume(BaseModel):
    name:str | None = None
    email:str | None = None
    phone:str | None = None

    total_experience_years: float | None = None

    skills: list[str] =[]
    experience: list[Experience] =[]
    education: list[str] =[]
    projects: list[str] =[]
    certifications: list[str] =[]

resume_schema = Resume.model_json_schema()

def final_score(job, resume):
    matchSchema = MatchResult.model_json_schema()
    prompt=f"""
        you are a HR Recruiter.

        Compare the candidates resume with the job description.

        JOB DESCRIPTION : {job.model_dump_json(indent=2)}

        CANDIDATE RESUME : {resume.model_dump_json(indent=2)}

        Return JSON Matching this schema: {matchSchema}

        Give me:
        1. Candidate Name
        2. Matching Skills
        3. Missing Important Skills
        4. Whether experince requirement is met
        5. Overall match percentage from 0 to 100
        6. A short final verdict

        Keep the response concise and easy to read.
    """ 

    message={
        "role":"user",
        "content":prompt
    }

    messages=[message]
    response_format={
        "type":"json_object"
    }

    response = client.chat.completions.create(model= model, messages = messages, response_format = response_format)
    data = json.loads(response.choices[0].message.content)
    return MatchResult(**data)



def parse_resume(resume_text):
    system_prompt=f"""
        You are an expert resume parser, 

        Extract information from the resume based on its meaning,
        not only based on exact section headings.

        Different resumes may use different headings.

        For example:
        - Experience
        - Professional Experience
        - Work History
        - Employment
        - Internships

        These may all conatin relevant experience.

        Skills may also appear in the skills section, work experience, internships or projects.

        Return ONLY valid JSON matching this schema:

        {resume_schema}

        Important rules:
        1. Do not invent information
        2. If a value iss not avaiable, return null
        3. If a list has no information, return an empty list
        4. Include internships inside experiences.
        5. Extract skills mentioned across the entire resume.
    """

    user_prompt= f"""
        Parse the following Resume:

        {resume_text}
    """

    message_system = {
        "role":"system",
        "content":system_prompt
    }

    message_user={
        "role":"user",
        "content":user_prompt
    }

    messages=[message_system, message_user]
    response_format={
        "type":"json_object"
    }

    response = client.chat.completions.create(model=model, messages= messages, response_format = response_format)

    raw_output = response.choices[0].message.content
    data = json.loads(raw_output)
    resume = Resume(**data)
    return resume


from pypdf import PdfReader
from docx import Document
def read_pdf(file_path):
    reader = PdfReader(file_path)
    text=""

    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text


def read_docx(file_path):
    document = Document(file_path)
    text=""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text+= paragraph.text + "\n"

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"

    return text

def read_resume(file_path):
    if file_path.suffix.lower() == ".pdf":
        return read_pdf(file_path)
    elif file_path.suffix.lower() == ".docx":
        return read_docx(file_path)
    else:
        return None

# Lets do it now
resume_folder = Path("Resumes")
all_results=[]
for file_path in resume_folder.iterdir():
    if file_path.suffix.lower() not in [".pdf", ".docx"]:
        continue

    print("\nProcessing:", file_path.name)
    resume_text = read_resume(file_path)
    parsed_resume = parse_resume(resume_text)

    time.sleep(5)
    result = final_score(job, parsed_resume)
    time.sleep(5)

    print("Score: ", result.score)

    all_results.append({
        "name":parsed_resume.name,
        "score":result.score,
        "details":result.details
    })

all_results.sort(
    key=lambda candidate: candidate["score"],
    reverse=True
)

top_2 = all_results[:2]
worst_2 = all_results[-2:]

print("TOP 2 Candidates")
for candidates in top_2:
    print(candidates["name"], "-", candidates["score"], "%")

    print(candidates["details"])


print("Worst 2 Candidates")
for candidates in worst_2:
    print(candidates["name"], "-", candidates["score"], "%")

    print(candidates["details"])